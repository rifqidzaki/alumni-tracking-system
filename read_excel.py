import openpyxl
wb = openpyxl.load_workbook(r"D:\SEMESTER 6\Rekayasa Kebutuhan\Alumni 2000-2025.xlsx", data_only=True)
ws = wb["Sheet1"]

print(f"Rows: {ws.max_row}, Cols: {ws.max_column}")
# Headers
for c in range(1, ws.max_column+1):
    print(f"  Col{c}: {ws.cell(1,c).value}")

# 5 sample rows
print("\nSAMPLE:")
for r in range(2,7):
    row = [str(ws.cell(r,c).value or "") for c in range(1,ws.max_column+1)]
    print(f"  {row}")
